"""Genera las hojas de datos (data sheets) en formato Excel.

Plantillas genéricas sin formato corporativo. La fase 2 agregará
encabezado, portada y código de documento con formato.

Uso: python scripts/generar_hojas_datos.py
Salida: content/07_Hojas_de_Datos/*.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

SALIDA = Path("content/07_Hojas_de_Datos")

FUENTE_TITULO = Font(bold=True, size=13)
FUENTE_SECCION = Font(bold=True, size=11)
FUENTE_ENCABEZADO = Font(bold=True)

COLUMNAS = ["Campo", "Valor", "Unidad", "Nota / Referencia"]
ANCHOS = {"A": 42, "B": 22, "C": 14, "D": 48}

# ---------------------------------------------------------------------------
# Definición de cada data sheet: (nombre_hoja, secciones)
# Cada sección: (título, [(campo, valor, unidad, nota), ...])
# Valor None -> celda vacía (plantilla genérica).
# ---------------------------------------------------------------------------

EASIHEAT = (
    "DS-CAL-001",
    "Hoja de Datos — Calentador Instantáneo de Agua a Vapor (tipo EasiHeat)",
    [
        ("Identificación del documento (fase 2: portada y encabezado)", [
            ("Código de documento", "DS-CAL-001", "-", "Reservado para formato corporativo"),
            ("Revisión", "0", "-", ""),
            ("Fecha", None, "-", ""),
            ("Proyecto", None, "-", ""),
            ("Preparó / Revisó / Aprobó", None, "-", ""),
        ]),
        ("1. Identificación del equipo", [
            ("Tag del equipo", "EWH-001 (ejemplo)", "-", "Etiqueta en P&ID"),
            ("Servicio", "Agua caliente sanitaria / proceso (ejemplo)", "-", ""),
            ("Ubicación", None, "-", "Planta / área"),
            ("Fabricante / Modelo", "Spirax Sarco EasiHeat (ejemplo)", "-", "Ver TI-P481-02"),
            ("Tipo", "Instantáneo, intercambiador de placas (ejemplo)", "-", "Sin tanque de almacenamiento"),
        ]),
        ("2. Condiciones de diseño — lado agua", [
            ("Fluido", "Agua", "-", ""),
            ("Temperatura de entrada", 25, "°C", "Ejemplo del caso del proyecto"),
            ("Temperatura de salida", 65, "°C", "Ejemplo del caso del proyecto"),
            ("Caudal de agua de diseño", None, "GPM / m³/h", "Según demanda del proceso"),
            ("Presión de agua de alimentación", None, "psig / barG", ""),
            ("Caída de presión permitida", None, "psi / bar", ""),
        ]),
        ("3. Condiciones de diseño — lado vapor", [
            ("Presión de vapor disponible (agotamiento PRS)", 25, "psig", "Ejemplo: salida de estación DS-PRS-001"),
            ("Temperatura de condensación aprox.", 130, "°C", "Vapor saturado a 25 psig"),
            ("Carga térmica de diseño (Q = m·cp·ΔT)", None, "kW / BTU/h", "Calcular con caudal y ΔT"),
            ("Consumo de vapor estimado", None, "kg/h / lb/h", "Q / calor latente a P de vapor"),
            ("Presión máxima admisible (PMA)", None, "psig / barG", "Según datasheet del fabricante"),
            ("Temperatura máxima admisible (TMA)", None, "°C", "Según datasheet del fabricante"),
        ]),
        ("4. Construcción y materiales", [
            ("Tipo de intercambiador", "Placas (ejemplo)", "-", ""),
            ("Material placas / tubos", None, "-", "Típico: acero inoxidable 316"),
            ("Material carcaza / bastidor", None, "-", ""),
            ("Empaques", None, "-", ""),
            ("Suministro", "Skid ensamblado y probado a presión (ejemplo)", "-", ""),
        ]),
        ("5. Control e instrumentación", [
            ("Sistema de control", "SIMS pantalla táctil (ejemplo)", "-", "O válvula modulante + controlador"),
            ("Válvula de control de vapor", None, "-", "Tipo, tamaño, Cv"),
            ("Sensor de temperatura", None, "-", "Rango y ubicación"),
            ("Mezclador termostático (si aplica)", None, "-", "Obligatorio en usos con contacto humano"),
            ("Alimentación eléctrica", None, "V / fases / Hz", ""),
        ]),
        ("6. Conexiones y suministros", [
            ("Conexión entrada vapor", None, "in / DN", ""),
            ("Conexión salida condensado", None, "in / DN", ""),
            ("Conexión entrada agua fría", None, "in / DN", ""),
            ("Conexión salida agua caliente", None, "in / DN", ""),
            ("Trampa de condensado", None, "-", "Flotador-termostática recomendada, FS 2-3"),
        ]),
        ("7. Notas y observaciones", [
            ("Nota 1", "Verificar capacidad contra el datasheet del modelo específico.", "-", ""),
            ("Nota 2", "Considerar condición de stall con cargas variables bajas; evaluar trampa-bomba.", "-", "Ref. cap. 13 Spirax"),
            ("Nota 3", "En agua dura evaluar tratamiento para evitar incrustaciones.", "-", ""),
            ("Nota 4", "Preferir instantáneo/semi-instantáneo para mitigar Legionella en DHW.", "-", ""),
        ]),
    ],
)

PRS = (
    "DS-PRS-001",
    "Hoja de Datos — Estación Reductora de Presión de Vapor (PRS)",
    [
        ("Identificación del documento (fase 2: portada y encabezado)", [
            ("Código de documento", "DS-PRS-001", "-", "Reservado para formato corporativo"),
            ("Revisión", "0", "-", ""),
            ("Fecha", None, "-", ""),
            ("Proyecto", None, "-", ""),
            ("Preparó / Revisó / Aprobó", None, "-", ""),
        ]),
        ("1. Identificación de la estación", [
            ("Tag de la estación", "PRS-001 (ejemplo)", "-", "Etiqueta en P&ID"),
            ("Servicio", "Reducción de presión de vapor para calentador (ejemplo)", "-", ""),
            ("Ubicación", None, "-", "Planta / área"),
            ("Equipo(s) alimentado(s)", "EWH-001 (ejemplo)", "-", "Ver DS-CAL-001"),
        ]),
        ("2. Condiciones de diseño", [
            ("Fluido", "Vapor saturado", "-", ""),
            ("Presión de entrada (caldera)", 150, "psig", "Ejemplo del caso del proyecto"),
            ("Presión de salida (setpoint)", 25, "psig", "Ejemplo del caso del proyecto"),
            ("Relación de reducción", "6:1 (ejemplo)", "-", "1 etapa hasta ~10:1; mayor → 2 etapas"),
            ("Flujo de vapor máximo", None, "kg/h / lb/h", "Demanda máxima del proceso — NUNCA dimensionar por Ø tubería"),
            ("Flujo de vapor mínimo", None, "kg/h / lb/h", "Si turndown alto → evaluar PRV en paralelo (lead/lag)"),
            ("Temperatura de entrada", None, "°C", "Saturación a P de entrada"),
        ]),
        ("3. Válvula reductora (PRV)", [
            ("Tipo de PRV", "Pilotada auto-actuada (ejemplo)", "-", "Directa: solo cargas pequeñas y constantes"),
            ("Modelo sugerido", "Spirax 25P / TLV COS-16 (ejemplo)", "-", "COSPECT integra filtro+separador+trampa"),
            ("Tamaño", None, "in / DN", "Típicamente 1-2 tamaños menor que la tubería"),
            ("Cv / Kv requerido", None, "-", "Calcular con presiones reales, no tabla genérica"),
            ("Clase / rating", None, "-", "Ej. ANSI 150/300, PN25/40"),
            ("Material cuerpo / internos", None, "-", "Hierro dúctil o acero; obturador Stellite"),
            ("Precisión de control", None, "±% setpoint", "Pilotada: alta; directa: ±10-20%"),
            ("Línea de sensado", "Sí, al lado de baja (ejemplo)", "-", "Tubería recta, ≥10 diámetros de la PRV"),
        ]),
        ("4. Válvula de seguridad (lado de baja — obligatoria)", [
            ("Tag", "PSV-001 (ejemplo)", "-", ""),
            ("Presión de tarado (set)", None, "psig", "10% o ~3-5 psi sobre el setpoint de baja"),
            ("Capacidad requerida", None, "kg/h / lb/h", "100% del flujo con la PRV fallada abierta"),
            ("Tamaño orificio / conexiones", None, "in / DN", "Según ASME / fabricante"),
            ("Detector de pérdidas (Spira-tec)", "Opcional", "-", ""),
        ]),
        ("5. Componentes de la estación (en orden de flujo)", [
            ("1. Válvula de aislamiento entrada", None, "in / DN", "Globo o compuerta"),
            ("2. Separador de humedad", None, "in / DN", "Vapor seco protege el asiento de la PRV"),
            ("2a. Trampa del separador", None, "-", "Flotador-termostática"),
            ("3. Filtro Y (strainer)", None, "in / DN", "Con pantalla; drenaje hacia abajo o lateral"),
            ("4. Manómetro lado alta", None, "rango", "Rango ~2× presión de entrada"),
            ("5. PRV (ver sección 3)", None, "-", ""),
            ("6. Manómetro lado baja", None, "rango", "Rango ~2× setpoint de baja"),
            ("7. Válvula de seguridad (ver sección 4)", None, "-", "Obligatoria: ASME B31.1 / PED"),
            ("8. Válvula de aislamiento salida", None, "in / DN", ""),
            ("Opcional: bypass manual", None, "in / DN", "Globo, para operación durante mantenimiento"),
            ("Opcional: check valve salida", None, "in / DN", ""),
        ]),
        ("6. Notas y observaciones", [
            ("Nota 1", "Sobredimensionar la PRV causa hunting y desgaste del asiento.", "-", ""),
            ("Nota 2", "Arranque lento: abrir el vapor de golpe causa golpe de ariete.", "-", ""),
            ("Nota 3", "Todo punto bajo aguas arriba requiere estación de trampa.", "-", ""),
            ("Nota 4", "La válvula modulante de temperatura del equipo va DESPUÉS de la PRS.", "-", ""),
            ("Nota 5", "Ref. interna: nota EstacionesReductoras, manual HookUp págs. 19-20.", "-", ""),
        ]),
    ],
)


def generar(archivo: str, codigo: str, titulo: str, secciones) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = codigo
    for col, ancho in ANCHOS.items():
        ws.column_dimensions[col].width = ancho

    fila = 1
    ws.cell(row=fila, column=1, value=titulo).font = FUENTE_TITULO
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    fila += 2

    for titulo_seccion, campos in secciones:
        celda = ws.cell(row=fila, column=1, value=titulo_seccion)
        celda.font = FUENTE_SECCION
        fila += 1
        for i, encabezado in enumerate(COLUMNAS, start=1):
            ws.cell(row=fila, column=i, value=encabezado).font = FUENTE_ENCABEZADO
        fila += 1
        for campo, valor, unidad, nota in campos:
            ws.cell(row=fila, column=1, value=campo)
            ws.cell(row=fila, column=2, value=valor)
            ws.cell(row=fila, column=3, value=unidad)
            celda_nota = ws.cell(row=fila, column=4, value=nota)
            celda_nota.alignment = Alignment(wrap_text=True, vertical="top")
            fila += 1
        fila += 1

    ws.freeze_panes = "A2"
    ruta = SALIDA / archivo
    wb.save(ruta)
    return ruta


def main():
    SALIDA.mkdir(parents=True, exist_ok=True)
    for args, archivo in [
        (EASIHEAT, "DS-CAL-001_Calentador_EasiHeat_Plantilla.xlsx"),
        (PRS, "DS-PRS-001_Estacion_Reductora_Plantilla.xlsx"),
    ]:
        ruta = generar(archivo, *args)
        print(f"Generado: {ruta}")


if __name__ == "__main__":
    main()
